import io
import zipfile

import pymupdf
import pytest
from docx import Document as WordDocument

from research_assist.application.documents import DocumentService
from research_assist.config import Settings
from research_assist.domain.errors import CapacityError, DocumentError
from research_assist.domain.models import Document, SourceSection
from research_assist.infrastructure.documents import chunk_sections, extract, safe_filename
from research_assist.infrastructure.embeddings import NoEmbeddings
from research_assist.infrastructure.repository import Repository


@pytest.mark.parametrize(
    "name,expected",
    [
        ("../../paper.pdf", "paper.pdf"),
        ("C:\\tmp\\study.pdf", "study.pdf"),
        ("Nghiên cứu.txt", "Nghiên cứu.txt"),
        ("file\x00.pdf", "file.pdf"),
    ],
)
def test_safe_filename(name: str, expected: str) -> None:
    assert safe_filename(name) == expected


@pytest.mark.parametrize("name", ["", "..", "/", "x" * 200])
def test_invalid_filename(name: str) -> None:
    with pytest.raises(DocumentError):
        safe_filename(name)


@pytest.mark.parametrize(
    "data,name",
    [(b"fake", "fake.pdf"), (b"abc", "data.exe"), (b"\xff\xfe\xff", "text.txt"), (b"", "empty.md")],
)
def test_invalid_document(data: bytes, name: str, settings: Settings) -> None:
    with pytest.raises(DocumentError):
        extract(data, name, settings)


def test_pdf_page_provenance(settings: Settings, document: Document) -> None:
    with pymupdf.open() as pdf:
        pdf.new_page().insert_text(
            (40, 50), "RAG combines retrieval and generation to use external evidence."
        )
        pdf.new_page().insert_text(
            (40, 50), "Self-RAG uses reflection tokens to critique generated answers."
        )
        data = pdf.tobytes()
    sections = extract(data, "paper.pdf", settings)
    assert [s.page for s in sections] == [1, 2]
    chunks = chunk_sections(sections, document, settings)
    assert [(c.page, c.start) for c in chunks] == [(1, 0), (2, 0)]
    assert "Self-RAG" not in chunks[0].text


def test_scan_pdf_explicitly_rejected(settings: Settings) -> None:
    with pymupdf.open() as pdf:
        pdf.new_page()
        data = pdf.tobytes()
    with pytest.raises(DocumentError, match="OCR"):
        extract(data, "scan.pdf", settings)


def test_pdf_page_limit(settings: Settings) -> None:
    settings.max_pages = 1
    with pymupdf.open() as pdf:
        pdf.new_page()
        pdf.new_page()
        data = pdf.tobytes()
    with pytest.raises(DocumentError, match="trang"):
        extract(data, "paper.pdf", settings)


def test_docx_preserves_table_flow_and_does_not_invent_pages(settings: Settings) -> None:
    doc = WordDocument()
    doc.add_paragraph("First paragraph: methods and data are clearly described.")
    doc.add_table(rows=1, cols=1).cell(0, 0).text = "Middle table: fifty documents."
    doc.add_paragraph("Last paragraph: evidence is evaluated conservatively.")
    stream = io.BytesIO()
    doc.save(stream)
    sections = extract(stream.getvalue(), "paper.docx", settings)
    assert sections[0].page is None
    assert (
        sections[0].text.index("First")
        < sections[0].text.index("Middle")
        < sections[0].text.index("Last")
    )


def test_zip_expansion_limit(settings: Settings) -> None:
    stream = io.BytesIO()
    with zipfile.ZipFile(stream, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("word/document.xml", b"x" * (25 * 1024 * 1024))
    with pytest.raises(DocumentError, match="giải nén"):
        extract(stream.getvalue(), "bomb.docx", settings)


def test_chunks_have_exact_offsets(document: Document, settings: Settings) -> None:
    settings.chunk_chars = 300
    settings.chunk_overlap = 40
    text = "Sentence with methods and evidence. " * 60
    chunks = chunk_sections([SourceSection(text=text, page=3)], document, settings)
    assert len(chunks) > 2
    for chunk in chunks:
        assert text[chunk.start : chunk.end] == chunk.text
        assert len(chunk.text) <= 300
        assert chunk.page == 3
    assert chunks[-1].end == len(text.rstrip())


def test_numbered_table_rows_remain_indexed(document: Document, settings: Settings) -> None:
    text = (
        "6. Results\n"
        "The evaluation uses three sets of sentences.\n"
        "1 General language 50 12.8\n"
        "2 Sports 25 24\n"
        "3 Tourism 25 26.5\n"
        "The groups were evaluated separately."
    )
    chunks = chunk_sections([SourceSection(text=text, page=12)], document, settings)
    assert any("2 Sports 25 24" in chunk.text for chunk in chunks)
    assert any("1 General language 50" in chunk.text for chunk in chunks)
    assert any("3 Tourism 25" in chunk.text for chunk in chunks)
    for chunk in chunks:
        assert chunk.text == text[chunk.start : chunk.end]


async def test_duplicate_upload_is_idempotent(settings: Settings) -> None:
    repository = Repository(settings.data_dir)
    service = DocumentService(repository, NoEmbeddings(), settings)
    data = b"RAG combines retrieval and generation with source evidence."
    first, duplicate = await service.ingest(data, "paper.txt")
    second, duplicate2 = await service.ingest(data, "renamed.txt")
    assert not duplicate and duplicate2
    assert first.id == second.id
    assert len(repository.documents()) == 1
    assert repository.chunk_count() == first.chunk_count
    assert repository.blob(first.id).read_bytes() == data


async def test_capacity_guard_does_not_leave_partial_document(settings: Settings) -> None:
    settings.max_chunks = 1
    settings.chunk_chars = 300
    service = DocumentService(Repository(settings.data_dir), NoEmbeddings(), settings)
    with pytest.raises(CapacityError):
        await service.ingest(b"Methods and retrieval evidence. " * 100, "large.txt")
    assert not service.repository.documents()
    assert not list(service.repository.blobs.iterdir())


async def test_rebuild_rejects_concurrent_corpus_change_without_overwriting(settings):
    repo = Repository(settings.data_dir)
    service = DocumentService(repo, NoEmbeddings(), settings)
    original, _ = await service.ingest(
        b"Original RAG evidence remains unchanged during rebuild.", "old.txt"
    )
    before = [c.id for c in repo.chunks([original.id])]
    other = DocumentService(Repository(settings.data_dir), NoEmbeddings(), settings)

    class ConcurrentUpload:
        fingerprint = None

        async def embed(self, texts, *, query=False):
            await other.ingest(
                b"Concurrent new RAG evidence must not be lost during rebuild.", "new.txt"
            )
            return []

    service.embedder = ConcurrentUpload()
    with pytest.raises(DocumentError, match="đã thay đổi"):
        await service.rebuild()
    assert [c.id for c in repo.chunks([original.id])] == before
    assert len(repo.documents()) == 2
